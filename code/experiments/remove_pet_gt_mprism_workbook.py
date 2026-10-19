"""Remove PET, GT and MPrISM from SSM_dataset_metric_matrix.xlsx: delete their rows in 5 sheets (shifting merged cells,
row heights and autofilters), drop the now-empty 交叉路线 column of 指标所需数据, shrink the COUNTIF ranges 6:17 -> 6:14,
fix the condition lists that name them and add a note in 说明.
usage: python -I remove_pet_gt_mprism_workbook.py <xlsx>"""
import re
import sys

import openpyxl
from openpyxl.worksheet.cell_range import CellRange

sys.stdout.reconfigure(encoding='utf-8')
path = sys.argv[1]
NAMES = {'PET', 'GT', 'MPrISM'}
NOTE = ('10. 2026-10-16 删掉了 PET、GT、MPrISM（本来就没实现，理由详见 docs/ssm_computation_steps.md 第 7 节）：'
        'PET、GT 是路口、横穿冲突的指标，要两方的路线交叉；自车视角的 20 秒片段里这种完整的相遇很少（demo 3 段一共只有 1 次），'
        '而同向跟车时 PET 就变成跟车时距，和 THW 重复。MPrISM 每帧都要解最坏情况的优化问题，计算量大、要定的参数多，'
        '原文预测时长只有 1 秒，普通场景里大多没有值。现在剩 9 个 SSM。')
TEXT_EDITS = {   # exact cell contents in △条件汇总
    '2D-TTC、ACT、MPrISM': '2D-TTC、ACT',
    'PET、GT': '（原来受影响的 PET、GT 已删除，这一条目前不影响任何指标）',
    'PET、GT（以及整个场景的最小 TTC）': '整个场景的最小 TTC（原来受影响的 PET、GT 已删除）',
}

wb = openpyxl.load_workbook(path)


def counts(wb):
    n_c = n_f = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                n_c += c.comment is not None
                n_f += isinstance(c.value, str) and c.value.startswith('=')
    return n_c, n_f


def delete_row(ws, r):
    max_row = ws.max_row
    merged = [CellRange(str(m)) for m in ws.merged_cells.ranges]
    for m in merged:
        ws.unmerge_cells(str(m))
    heights = {i: ws.row_dimensions[i].height for i in range(r, max_row + 1)}
    ws.delete_rows(r)
    for i in range(r, max_row):
        ws.row_dimensions[i].height = heights.get(i + 1)
    for m in merged:
        if m.min_row > r:
            m.shift(row_shift=-1)
        elif m.max_row >= r:
            raise RuntimeError(f'merged range {m} spans deleted row {r}')
        ws.merge_cells(m.coord)
    if ws.auto_filter.ref:
        cr = CellRange(ws.auto_filter.ref)
        if cr.max_row >= r:
            cr.expand(down=-1)
            ws.auto_filter.ref = cr.coord


before = counts(wb)
# 1. texts first
n_edits = 0
for row in wb['△条件汇总'].iter_rows():
    for c in row:
        if isinstance(c.value, str) and c.value in TEXT_EDITS:
            print(f'△条件汇总!{c.coordinate}: {c.value} -> {TEXT_EDITS[c.value]}')
            c.value = TEXT_EDITS[c.value]
            n_edits += 1
assert n_edits == 3, n_edits
ws = wb['说明']
assert ws['B20'].value.startswith('9. ') and ws['B21'].value is None
ws['B21'].value = NOTE
ws['B21'].alignment = ws['B20'].alignment.copy()
ws['B21'].font = ws['B20'].font.copy()
if ws.row_dimensions[20].height:
    ws.row_dimensions[21].height = ws.row_dimensions[20].height

# 2. delete rows named after the metrics
plan = {'按指标类型分类': 'B', '按碰撞类型分类': 'C', '指标×数据集': 'A', '条件明细': 'B', '指标所需数据': 'A'}
lost = n_matrix_rows = total = 0
for name, col in plan.items():
    ws = wb[name]
    rows = [c.row for c in ws[col] if c.value in NAMES]
    lost += sum(1 for r in rows for c in ws[r] if c.comment is not None)
    for r in sorted(rows, reverse=True):
        delete_row(ws, r)
    if name == '指标×数据集':
        n_matrix_rows = len(rows)
    total += len(rows)
    print(f'{name}: deleted {len(rows)} rows {rows if len(rows) < 12 else str(rows[:6]) + "..."}')
print('rows deleted in total:', total)

# 3. 指标所需数据: the 交叉路线 column (H) is empty now -> drop it and shift the column widths
ws = wb['指标所需数据']
assert ws['H1'].value == '交叉路线', ws['H1'].value
used = [c.coordinate for c in ws['H'][1:] if c.value not in (None, '')]
assert not used, f'交叉路线 still used in {used}'
widths = {k: ws.column_dimensions[k].width for k in ('I', 'J') if k in ws.column_dimensions}
ws.delete_cols(8)
ws.column_dimensions['H'].width = widths.get('I')
ws.column_dimensions['I'].width = widths.get('J')
if 'J' in ws.column_dimensions:
    del ws.column_dimensions['J']
print('指标所需数据 header now:', [c.value for c in ws[1]])

# 4. COUNTIF ranges over the metric rows: 6:17 -> 6:14
old_last, new_last = 17, 17 - n_matrix_rows
pat = re.compile(rf"(\$?[A-Z]{{1,3}}\$?)6:(\$?[A-Z]{{1,3}}\$?){old_last}\b")
fixed = 0
for name in ('指标×数据集', '自车图像数据集对比'):
    for row in wb[name].iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('=') and pat.search(c.value):
                c.value = pat.sub(rf'\g<1>6:\g<2>{new_last}', c.value)
                fixed += 1
print(f'formulas re-ranged 6:{old_last} -> 6:{new_last}:', fixed)

wb.save(path)
wb2 = openpyxl.load_workbook(path)
after = counts(wb2)
print('comments before/after:', before[0], after[0], '(expected loss', lost, ') | formulas:', before[1], after[1])
word = re.compile(r'\bPET\b|\bGT\b|MPrISM|MPrTTC|交叉路线')
left = [(ws.title, c.coordinate, c.value[:60]) for ws in wb2.worksheets for row in ws.iter_rows() for c in row
        if isinstance(c.value, str) and word.search(c.value)]
print('mentions left:', left)
ws = wb2['指标×数据集']
print('metric rows now:', [ws.cell(r, 1).value for r in range(6, new_last + 1)])
print('summary rows:', [(r, ws.cell(r, 1).value, ws.cell(r, 5).value) for r in range(new_last + 1, new_last + 7)])
for name in ('按指标类型分类', '按碰撞类型分类', '条件明细', '指标所需数据'):
    print(name, 'autofilter', wb2[name].auto_filter.ref, 'max_row', wb2[name].max_row)
odd = [c.value for row in wb2['自车图像数据集对比'].iter_rows() for c in row
       if isinstance(c.value, str) and c.value.startswith('=') and f'6:{old_last}' in c.value]
print(f'自车图像数据集对比 formulas still on 6:{old_last}:', odd)
