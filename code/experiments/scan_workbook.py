"""Read-only scan: where does a metric appear in SSM_dataset_metric_matrix.xlsx?
usage: python -I scan_workbook.py <xlsx> <regex>"""
import re
import sys

import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
path, pattern = sys.argv[1], sys.argv[2]
pat = re.compile(pattern, re.I)
wb = openpyxl.load_workbook(path)
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and pat.search(c.value):
                print(f'{ws.title}!{c.coordinate}: {c.value[:160]!r}')
            if c.comment is not None and pat.search(c.comment.text):
                print(f'  [comment] {ws.title}!{c.coordinate}: {c.comment.text[:160]!r}')
plan = {'按指标类型分类': 'B', '按碰撞类型分类': 'C', '指标×数据集': 'A', '条件明细': 'B', '指标所需数据': 'A'}
for name, col in plan.items():
    ws = wb[name]
    vals = [c.value for c in ws[col] if c.value is not None]
    print(name, col, 'distinct:', sorted(set(map(str, vals)))[:60])
